"""Phase 5: ZPL, PDF/image overlay and DOCX templates each render a test label (no Chromium needed)."""
import base64
import io
import json
import zipfile

from PIL import Image
from pypdf import PdfReader

from tests.conftest import login
from tests.helpers import body_for, pdf_text, soffice, squash  # noqa: F401  (soffice is a fixture)


def make_active(api, content, fname, name, **data):
    r = api.post("/api/templates", files={"file": (fname, content)}, data={"name": name, "size": "A6", **data})
    assert r.status_code == 201, r.text
    t = r.json()
    a = api.post(f"/api/templates/{t['id']}/activate")
    assert a.status_code == 200, a.text
    return t["id"]


ZPL = b"^XA\n^FO20,20^FD{{so_number}}^FS\n^FO20,60^FD{{description}}^FS\n^FO20,100^BEN,60^FD{{barcode}}^FS\n^FO20,200^FDPCS {{pcs}}^FS\n^XZ"


def test_zpl_template_renders_substituted_text_repeated_per_copy(api):
    login(api, "alice")
    tid = make_active(api, ZPL, "label.zpl", "ZPL label")
    api.odoo.data["product.product"][0]["name"] = "Name ^with~ commands"
    r = api.post("/api/render/print", json=body_for(api, tid, copies=2))
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/plain")
    text = r.text
    assert text.count("^XA") == 2
    assert "^FDS00123^FS" in text and "^FDPCS 10^FS" in text and "^FD5901234123457^FS" in text
    assert "^FDName  with  commands^FS" in text  # ^ and ~ from data can never inject ZPL commands
    assert r.headers["x-print-job-id"]


def test_zpl_can_be_sent_to_a_configured_network_printer(api, monkeypatch):
    from app.config import get_settings
    from app.render import router as render_router

    login(api, "alice")
    tid = make_active(api, ZPL, "label.zpl", "ZPL label")
    get_settings().printers["zebra1"] = "10.0.0.5:9100"
    sent = {}
    monkeypatch.setattr(render_router, "send_zpl", lambda target, data: sent.update(target=target, data=data))
    r = api.post("/api/render/print", json=body_for(api, tid, printer="zebra1", send_to_printer=True))
    assert r.headers["x-printer-status"] == "sent" and sent["target"] == "10.0.0.5:9100" and b"^FDS00123^FS" in sent["data"]
    r = api.post("/api/render/print", json=body_for(api, tid, printer="nope", send_to_printer=True))
    assert r.headers["x-printer-status"] == "unknown-printer"


def overlay_zip(bg_pdf: bool = False) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        if bg_pdf:
            from reportlab.pdfgen import canvas

            b = io.BytesIO()
            c = canvas.Canvas(b, pagesize=(105 * 72 / 25.4, 148 * 72 / 25.4))
            c.drawString(20, 20, "PREPRINTED FORM")
            c.save()
            z.writestr("background.pdf", b.getvalue())
        else:
            img = io.BytesIO()
            Image.new("RGB", (298, 420), "white").save(img, "PNG")
            z.writestr("background.png", img.getvalue())
        z.writestr("fields.json", json.dumps({"size": "A6", "fields": [
            {"placeholder": "so_number", "kind": "text", "x": 10, "y": 10, "w": 60, "size": 14, "font": "Helvetica-Bold"},
            {"placeholder": "description", "kind": "text", "x": 10, "y": 25, "w": 40, "h": 14, "size": 12, "overflow": "wrap"},
            {"placeholder": "item_code", "kind": "text", "x": 10, "y": 50, "w": 30, "size": 10, "overflow": "truncate"},
            {"placeholder": "barcode", "kind": "barcode", "x": 10, "y": 70, "w": 60, "h": 20},
            {"placeholder": "photo", "kind": "image", "x": 10, "y": 95, "w": 30, "h": 30, "optional": True},
            {"placeholder": "pcs", "kind": "text", "x": 60, "y": 100, "size": 12, "align": "left"}]}))
    return buf.getvalue()


def test_pdf_overlay_png_background_places_fields_at_coordinates(api):
    login(api, "alice")
    tid = make_active(api, overlay_zip(), "form.zip", "Overlay PNG")
    r = api.post("/api/render/preview", json=body_for(api, tid))
    assert r.status_code == 200, r.text
    pdf = base64.b64decode(r.json()["pdf_base64"])
    text = squash(pdf_text(pdf))
    assert "S00123" in text and "Thermometer" in text and "220-TD" in text and "10" in text
    page = PdfReader(io.BytesIO(pdf)).pages[0]
    assert len(page.images) >= 3  # background + barcode + product photo
    assert round(float(page.mediabox.width) / (72 / 25.4)) == 105


def test_pdf_overlay_pdf_background_keeps_the_preprinted_form(api):
    login(api, "alice")
    tid = make_active(api, overlay_zip(bg_pdf=True), "form.zip", "Overlay PDF")
    r = api.post("/api/render/preview", json=body_for(api, tid))
    text = squash(pdf_text(base64.b64decode(r.json()["pdf_base64"])))
    assert "PREPRINTED FORM" in text and "S00123" in text


def test_overlay_upload_needs_fields_json_with_background(api):
    login(api, "alice")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("fields.json", "{}")
    r = api.post("/api/templates", files={"file": ("x.zip", buf.getvalue())}, data={"name": "nobg"})
    assert r.status_code == 422


def make_docx() -> bytes:
    from docx import Document

    d = Document()
    d.add_paragraph("SO {{ so_number }} item {{ line.product.code }} pcs {{ pcs }}")
    d.add_paragraph("{% if pefc %}PEFC CERTIFIED{% endif %}")
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def test_docx_placeholders_are_scanned_and_mapped(api):
    login(api, "alice")
    r = api.post("/api/templates", files={"file": ("l.docx", make_docx())}, data={"name": "Word", "size": "A6"})
    assert r.status_code == 201, r.text
    maps = {m["placeholder"]: m["catalog_key"] for m in r.json()["versions"][0]["mappings"]}
    assert maps == {"so_number": "header.name", "line.product.code": "line.product.code", "pcs": "calc.pcs"}


def test_docx_renders_through_libreoffice(api, soffice):
    login(api, "alice")
    tid = make_active(api, make_docx(), "l.docx", "Word")
    r = api.post("/api/render/print", json=body_for(api, tid))
    assert r.status_code == 200, r.text
    assert "SO S00123 item 220-TD pcs 10" in squash(pdf_text(r.content))


def test_docx_template_cannot_escape_the_jinja_sandbox(api, soffice):
    from docx import Document

    d = Document()
    d.add_paragraph("{{ so_number }} {{ ''.__class__.__mro__[1].__subclasses__() }}")
    out = io.BytesIO()
    d.save(out)
    login(api, "alice")
    r = api.post("/api/templates", files={"file": ("e.docx", out.getvalue())}, data={"name": "Evil"})
    assert r.status_code == 201  # scan ignores it (not a plain placeholder)
    tid = r.json()["id"]
    api.post(f"/api/templates/{tid}/activate")
    res = api.post("/api/render/print", json=body_for(api, tid))
    assert res.status_code == 422 and "could not be rendered" in res.text  # blocked by the Jinja sandbox, not executed
